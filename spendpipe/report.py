"""Automated monthly operating report (Excel).

One command rebuilds the whole workbook from the SQLite database:

    python -m spendpipe report --month 2026-03

Derived numbers (shares, changes, averages, cumulative %, totals) are written
as live Excel formulas over the aggregate tables the database supplies, so the
workbook recalculates if anyone edits a cell.
"""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C

FONT = "Arial"
NAVY = "1F3864"
LIGHT = "DCE6F1"
GREY = "595959"
WARN_FILL = "FFF2CC"

F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=9, italic=True, color=GREY)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_NOTE = Font(name=FONT, size=9, color=GREY)
FILL_HDR = PatternFill("solid", start_color=NAVY)
FILL_TOTAL = PatternFill("solid", start_color=LIGHT)
FILL_WARN = PatternFill("solid", start_color=WARN_FILL)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FMT_INT = "#,##0;(#,##0);-"
FMT_USD = "$#,##0;($#,##0);-"
FMT_USDM = '$#,##0.0,,"M";($#,##0.0,,"M");-'
FMT_PCT = "0.0%;(0.0%);-"


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _hdr(ws, row: int, headers: list[str], col: int = 1) -> None:
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill, c.border = F_HDR, FILL_HDR, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _put(ws, row: int, col: int, value, fmt: str | None = None, bold: bool = False, fill=None, wrap: bool = False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F_BOLD if bold else F_BODY
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def _title(ws, text: str, sub: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _key(period: pd.Period) -> str:
    return period.strftime("%Y-%m")


def _month_bounds(month: str) -> tuple[str, str]:
    """(first day of T12M window, first day of the month after) as ISO strings."""
    p = pd.Period(month, "M")
    return (p - 11).start_time.strftime("%Y-%m-%d"), (p + 1).start_time.strftime("%Y-%m-%d")


def default_month(con: sqlite3.Connection) -> str:
    """Latest award month that ended at least PROVISIONAL_DAYS before the newest publication."""
    through = pd.to_datetime(con.execute("SELECT MAX(publication_date) FROM fact_contract_event").fetchone()[0])
    months = [r[0] for r in con.execute(
        "SELECT DISTINCT substr(award_month,1,7) FROM fact_contract WHERE award_month IS NOT NULL ORDER BY 1")]
    settled = [m for m in months
               if pd.Period(m, "M").end_time.normalize() + pd.Timedelta(days=C.PROVISIONAL_DAYS) <= through]
    return settled[-1] if settled else months[-2]


# --------------------------------------------------------------------------
# main entry point
# --------------------------------------------------------------------------
def generate_report(db_path: str, month: str | None = None, out_path: str | None = None) -> str:
    con = sqlite3.connect(db_path)
    q = lambda sql, params=(): pd.read_sql_query(sql, con, params=params)

    month = month or default_month(con)
    period = pd.Period(month, "M")
    prior_key, yago_key = _key(period - 1), _key(period - 12)
    t12_start, t12_end = _month_bounds(month)
    m_start, m_end = period.start_time.strftime("%Y-%m-%d"), (period + 1).start_time.strftime("%Y-%m-%d")

    through = con.execute("SELECT MAX(publication_date) FROM fact_contract_event").fetchone()[0]
    provisional = period.end_time.normalize() + pd.Timedelta(days=C.PROVISIONAL_DAYS) > pd.to_datetime(through)
    n_contracts_total = con.execute("SELECT COUNT(*) FROM fact_contract").fetchone()[0]
    lag = q("SELECT publication_lag_days d FROM fact_contract WHERE publication_lag_days IS NOT NULL")["d"]

    # ---- aggregate tables from the database ------------------------------
    monthly = q("""
        SELECT substr(award_month,1,7) AS m,
               COUNT(*) AS n,
               SUM(CASE WHEN final_value > 0 THEN 1 ELSE 0 END) AS nv,
               SUM(CASE WHEN final_value > 0 THEN final_value ELSE 0 END) AS v,
               SUM(CASE WHEN final_value > 0 AND final_value < 100000000 THEN final_value ELSE 0 END) AS vx,
               COUNT(DISTINCT supplier_id) AS ns,
               SUM(CASE WHEN final_value > 0 AND sourcing_type = 'Competitive' THEN final_value ELSE 0 END) AS comp,
               SUM(CASE WHEN final_value > 0 AND sourcing_type = 'Non-competitive' THEN final_value ELSE 0 END) AS noncomp
        FROM fact_contract WHERE award_month IS NOT NULL GROUP BY 1 ORDER BY 1""")

    cat = q("""
        SELECT c.category_name AS category,
               SUM(CASE WHEN f.award_month >= ? AND f.award_month < ? THEN f.final_value ELSE 0 END) AS month_v,
               SUM(f.final_value) AS t12_v,
               COUNT(*) AS t12_n,
               COUNT(DISTINCT f.supplier_id) AS t12_s
        FROM fact_contract f JOIN dim_category c ON c.category_id = f.category_id
        WHERE f.final_value > 0 AND f.award_month >= ? AND f.award_month < ?
        GROUP BY c.category_name ORDER BY t12_v DESC""", (m_start, m_end, t12_start, t12_end))

    sup = q("""
        SELECT s.supplier_name AS supplier, s.region AS region,
               COUNT(*) AS n, SUM(f.final_value) AS t12_v,
               SUM(CASE WHEN f.award_month >= ? AND f.award_month < ? THEN f.final_value ELSE 0 END) AS month_v
        FROM fact_contract f JOIN dim_supplier s ON s.supplier_id = f.supplier_id
        WHERE f.final_value > 0 AND f.award_month >= ? AND f.award_month < ?
        GROUP BY f.supplier_id ORDER BY t12_v DESC LIMIT 15""", (m_start, m_end, t12_start, t12_end))
    t12_total = con.execute(
        "SELECT COALESCE(SUM(final_value),0) FROM fact_contract WHERE final_value > 0 AND award_month >= ? AND award_month < ?",
        (t12_start, t12_end)).fetchone()[0]

    src = q("""
        SELECT sourcing_type AS t,
               SUM(CASE WHEN award_month >= ? AND award_month < ? THEN final_value ELSE 0 END) AS month_v,
               SUM(final_value) AS t12_v
        FROM fact_contract WHERE final_value > 0 AND award_month >= ? AND award_month < ?
        GROUP BY 1""", (m_start, m_end, t12_start, t12_end)).set_index("t")
    lim = q("""
        SELECT COALESCE(limited_tendering_reason, 'Not stated') AS reason, COUNT(*) AS n, SUM(final_value) AS v
        FROM fact_contract WHERE final_value > 0 AND sourcing_type = 'Non-competitive'
              AND award_month >= ? AND award_month < ?
        GROUP BY 1 ORDER BY v DESC LIMIT 6""", (t12_start, t12_end))

    amend_n = con.execute(
        "SELECT COUNT(*) FROM fact_contract WHERE n_amendments > 0 AND last_amendment_date >= ? AND last_amendment_date < ?",
        (m_start, m_end)).fetchone()[0]
    amend = q("""
        SELECT f.contract_number AS cn, COALESCE(f.title, '') AS title, s.supplier_name AS supplier,
               c.category_name AS category, f.original_amount AS orig, f.final_value AS final, f.n_amendments AS na
        FROM fact_contract f
        JOIN dim_supplier s ON s.supplier_id = f.supplier_id
        JOIN dim_category c ON c.category_id = f.category_id
        WHERE f.growth_eligible = 1 AND f.last_amendment_date >= ? AND f.last_amendment_date < ?
        ORDER BY (f.final_value - f.original_amount) DESC LIMIT 10""", (m_start, m_end))

    dq = q("SELECT issue_type, COUNT(*) AS n FROM dq_issue GROUP BY 1 ORDER BY n DESC")
    con.close()

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_cat = wb.create_sheet("Category")
    ws_sup = wb.create_sheet("Suppliers")
    ws_src = wb.create_sheet("Sourcing")
    ws_am = wb.create_sheet("Amendments")
    ws_mon = wb.create_sheet("Monthly Data")
    ws_dq = wb.create_sheet("Data Quality")
    ws_nt = wb.create_sheet("Notes")

    # ======================= Monthly Data =================================
    _title(ws_mon, "Monthly awards - all months in the file",
           "Aggregates from the database (contracts grouped by award month). Columns I-K are formulas.")
    mh = ["Month", "Contracts awarded", "Valued contracts (value > $0)", "Awarded value (CAD)",
          "Awarded value excl. contracts over $100M", "Distinct suppliers", "Competitive value",
          "Non-competitive value", "Avg valued contract", "Competitive share of reported value",
          "Contracts awarded - 3-mo avg"]
    _hdr(ws_mon, 4, mh)
    r0 = 5
    for i, row in monthly.iterrows():
        r = r0 + i
        _put(ws_mon, r, 1, row["m"])
        _put(ws_mon, r, 2, int(row["n"]), FMT_INT)
        _put(ws_mon, r, 3, int(row["nv"]), FMT_INT)
        _put(ws_mon, r, 4, float(row["v"]), FMT_USD)
        _put(ws_mon, r, 5, float(row["vx"]), FMT_USD)
        _put(ws_mon, r, 6, int(row["ns"]), FMT_INT)
        _put(ws_mon, r, 7, float(row["comp"]), FMT_USD)
        _put(ws_mon, r, 8, float(row["noncomp"]), FMT_USD)
        _put(ws_mon, r, 9, f"=IFERROR(D{r}/C{r},0)", FMT_USD)
        _put(ws_mon, r, 10, f'=IFERROR(G{r}/(G{r}+H{r}),"n/a")', FMT_PCT)
        _put(ws_mon, r, 11, f"=AVERAGE(B{max(r0, r - 2)}:B{r})", "#,##0.0")
    r_last = r0 + len(monthly) - 1
    _widths(ws_mon, {"A": 10, "B": 13, "C": 15, "D": 17, "E": 20, "F": 12, "G": 16, "H": 16, "I": 15, "J": 17, "K": 15})
    ws_mon.freeze_panes = "B5"

    last24_start = max(r0, r_last - 23)
    ch = LineChart()
    ch.title = "Contracts awarded per month (last 24 months)"
    ch.height, ch.width = 7.5, 16
    ch.add_data(Reference(ws_mon, min_col=2, min_row=last24_start, max_row=r_last), titles_from_data=False)
    ch.add_data(Reference(ws_mon, min_col=11, min_row=last24_start, max_row=r_last), titles_from_data=False)
    ch.set_categories(Reference(ws_mon, min_col=1, min_row=last24_start, max_row=r_last))
    ch.series[0].tx = SeriesLabel(v="Contracts awarded")
    ch.series[1].tx = SeriesLabel(v="3-month average")
    for series in ch.series:
        series.smooth = False
    ch.y_axis.title = "Contracts"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws_mon.add_chart(ch, "M4")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "Awarded value excl. contracts over $100M (CAD)"
    ch2.height, ch2.width = 7.5, 16
    ch2.add_data(Reference(ws_mon, min_col=5, min_row=last24_start, max_row=r_last), titles_from_data=False)
    ch2.set_categories(Reference(ws_mon, min_col=1, min_row=last24_start, max_row=r_last))
    ch2.series[0].tx = SeriesLabel(v="Awarded value excl. mega-contracts")
    ch2.legend = None
    ch2.x_axis.delete = False
    ch2.y_axis.delete = False
    ch2.y_axis.number_format = '$#,##0,,"M"'
    ws_mon.add_chart(ch2, "M20")

    mrng = lambda col: f"'Monthly Data'!${col}${r0}:${col}${r_last}"

    # ======================= Summary =====================================
    _title(ws_sum, f"Monthly Procurement Operating Report - {period.strftime('%B %Y')}",
           "Source: CanadaBuys contract history (federal contracts), Open Government Licence - Canada. "
           f"Data through publication date {through}. Amounts in CAD.")
    status = (f"PROVISIONAL: this month ended within {C.PROVISIONAL_DAYS} days of the newest publication date in the "
              "file, and awards are often posted months late, so figures may still rise."
              if provisional else
              f"Settled month: ended more than {C.PROVISIONAL_DAYS} days before the newest publication date "
              "(heuristic - late postings are still possible; see Notes).")
    ws_sum["A3"] = status
    ws_sum["A3"].font = Font(name=FONT, size=9, bold=provisional, color="7F6000" if provisional else GREY)
    if provisional:
        ws_sum["A3"].fill = FILL_WARN
    ws_sum["A4"] = "Month key"
    ws_sum["A4"].font = F_NOTE
    for col, k in ((2, month), (3, prior_key), (6, yago_key)):
        c = ws_sum.cell(row=4, column=col, value=k)
        c.font, c.alignment = F_NOTE, Alignment(horizontal="center")
    _hdr(ws_sum, 5, ["KPI (by award month)", "This month", "Prior month", "Change", "% change",
                     "Same month last year", "% vs last year"])

    def lookup(col: str, key_cell: str) -> str:
        return f'=IFERROR(INDEX({mrng(col)},MATCH({key_cell},{mrng("A")},0)),"n/a")'

    kpis = [
        ("Contracts awarded", "B", FMT_INT),
        ("  of which valued (value > $0)", "C", FMT_INT),
        ("Awarded value (CAD)", "D", FMT_USDM),
        ("Awarded value excl. contracts over $100M", "E", FMT_USDM),
        ("Distinct suppliers awarded", "F", FMT_INT),
        ("Average valued contract (CAD)", "I", FMT_USD),
        ("Competitive share of reported value", "J", FMT_PCT),
    ]
    for i, (label, col, fmt) in enumerate(kpis):
        r = 6 + i
        _put(ws_sum, r, 1, label, bold=not label.startswith("  "))
        _put(ws_sum, r, 2, lookup(col, "B$4"), fmt)
        _put(ws_sum, r, 3, lookup(col, "C$4"), fmt)
        _put(ws_sum, r, 4, f'=IFERROR(B{r}-C{r},"n/a")', fmt if col != "J" else FMT_PCT)
        _put(ws_sum, r, 5, f'=IFERROR(IF(C{r}=0,"n/a",B{r}/C{r}-1),"n/a")' if col != "J" else "", FMT_PCT)
        _put(ws_sum, r, 6, lookup(col, "F$4"), fmt)
        _put(ws_sum, r, 7, f'=IFERROR(IF(F{r}=0,"n/a",B{r}/F{r}-1),"n/a")' if col != "J" else "", FMT_PCT)
    r = 6 + len(kpis)
    _put(ws_sum, r, 1, "Contracts with $0 value (standing offers, arrangements)", bold=True)
    for col in (2, 3, 6):
        L = get_column_letter(col)
        _put(ws_sum, r, col, f'=IFERROR({L}6-{L}7,"n/a")', FMT_INT)
    _put(ws_sum, r, 4, f'=IFERROR(B{r}-C{r},"n/a")', FMT_INT)
    _put(ws_sum, r, 5, f'=IFERROR(IF(C{r}=0,"n/a",B{r}/C{r}-1),"n/a")', FMT_PCT)
    _put(ws_sum, r, 7, f'=IFERROR(IF(F{r}=0,"n/a",B{r}/F{r}-1),"n/a")', FMT_PCT)

    rr = r + 2
    ws_sum.cell(row=rr, column=1, value="Reading this report").font = F_BOLD
    notes = [
        "Awarded value = final cumulative value of contracts awarded in the month (valued contracts only), attributed to the award month.",
        "A single very large award can dominate a month; the 'excl. contracts over $100M' row shows the underlying run-rate.",
        "Competitive share = competitive value / (competitive + non-competitive value); contracts with no method reported are excluded. Its 'Change' is in percentage points.",
        "Detail tabs: Category, Suppliers, Sourcing (trailing 12 months ending this month), Amendments, Monthly Data, Data Quality, Notes.",
    ]
    for i, t in enumerate(notes):
        c = ws_sum.cell(row=rr + 1 + i, column=1, value=t)
        c.font = F_NOTE
    _widths(ws_sum, {"A": 52, "B": 16, "C": 16, "D": 16, "E": 11, "F": 20, "G": 14})
    ws_sum.freeze_panes = "A6"

    # ======================= Category =====================================
    _title(ws_cat, f"Spend by category - {period.strftime('%B %Y')} and trailing 12 months",
           f"T12M = 12 months ending {month}. Valued contracts only. Shares and averages are formulas.")
    _hdr(ws_cat, 4, ["Category", "This month (CAD)", "T12M (CAD)", "Share of T12M", "T12M monthly avg",
                     "This month vs avg", "Contracts T12M", "Suppliers T12M"])
    c0 = 5
    for i, row in cat.iterrows():
        r = c0 + i
        _put(ws_cat, r, 1, row["category"])
        _put(ws_cat, r, 2, float(row["month_v"]), FMT_USD)
        _put(ws_cat, r, 3, float(row["t12_v"]), FMT_USD)
        _put(ws_cat, r, 4, f"=IFERROR(C{r}/C${c0 + len(cat)},0)", FMT_PCT)
        _put(ws_cat, r, 5, f"=C{r}/12", FMT_USD)
        _put(ws_cat, r, 6, f'=IFERROR(B{r}/E{r}-1,"n/a")', FMT_PCT)
        _put(ws_cat, r, 7, int(row["t12_n"]), FMT_INT)
        _put(ws_cat, r, 8, int(row["t12_s"]), FMT_INT)
    rt = c0 + len(cat)
    _put(ws_cat, rt, 1, "Total", bold=True, fill=FILL_TOTAL)
    for col, L in ((2, "B"), (3, "C"), (5, "E"), (7, "G")):
        _put(ws_cat, rt, col, f"=SUM({L}{c0}:{L}{rt - 1})", FMT_USD if col != 7 else FMT_INT, bold=True, fill=FILL_TOTAL)
    _put(ws_cat, rt, 4, f"=SUM(D{c0}:D{rt - 1})", FMT_PCT, bold=True, fill=FILL_TOTAL)
    _put(ws_cat, rt, 6, f'=IFERROR(B{rt}/E{rt}-1,"n/a")', FMT_PCT, bold=True, fill=FILL_TOTAL)
    _put(ws_cat, rt, 8, "", bold=True, fill=FILL_TOTAL)
    ws_cat.cell(row=rt + 1, column=1,
                value="Categories come from the UNSPSC code (keyword rules for legacy records). Supplier counts are distinct within each category.").font = F_NOTE
    _widths(ws_cat, {"A": 44, "B": 17, "C": 18, "D": 13, "E": 17, "F": 14, "G": 13, "H": 13})
    ws_cat.freeze_panes = "B5"
    bc = BarChart()
    bc.type = "bar"
    bc.title = "T12M value by category (CAD)"
    bc.height, bc.width = 9, 18
    bc.add_data(Reference(ws_cat, min_col=3, min_row=c0, max_row=rt - 1), titles_from_data=False)
    bc.set_categories(Reference(ws_cat, min_col=1, min_row=c0, max_row=rt - 1))
    bc.series[0].tx = SeriesLabel(v="T12M value")
    bc.legend = None
    bc.x_axis.scaling.orientation = "maxMin"
    bc.x_axis.delete = False
    bc.y_axis.delete = False
    bc.y_axis.number_format = '$#,##0,,"M"'
    ws_cat.add_chart(bc, f"A{rt + 4}")

    # ======================= Suppliers ====================================
    _title(ws_sup, f"Top 15 suppliers - trailing 12 months ending {month}",
           "Supplier names are merged with rule-based matching (no fuzzy matching), so some related entities may still appear separately.")
    _put(ws_sup, 3, 1, "Total T12M valued awards, all suppliers (CAD)", bold=True)
    _put(ws_sup, 3, 2, float(t12_total), FMT_USD, bold=True)
    ws_sup.merge_cells("B3:C3")
    _hdr(ws_sup, 5, ["Rank", "Supplier", "Region", "Contracts T12M", "T12M value (CAD)", "Share of T12M",
                     "Cumulative share", "This month (CAD)"])
    s0 = 6
    for i, row in sup.iterrows():
        r = s0 + i
        _put(ws_sup, r, 1, i + 1, "0")
        _put(ws_sup, r, 2, row["supplier"], wrap=False)
        _put(ws_sup, r, 3, row["region"])
        _put(ws_sup, r, 4, int(row["n"]), FMT_INT)
        _put(ws_sup, r, 5, float(row["t12_v"]), FMT_USD)
        _put(ws_sup, r, 6, f"=IFERROR(E{r}/$B$3,0)", FMT_PCT)
        _put(ws_sup, r, 7, f"=F{r}" if i == 0 else f"=G{r - 1}+F{r}", FMT_PCT)
        _put(ws_sup, r, 8, float(row["month_v"]), FMT_USD)
    if len(sup):
        rs = s0 + len(sup)
        _put(ws_sup, rs, 2, "Top 10 share of T12M value", bold=True, fill=FILL_TOTAL)
        _put(ws_sup, rs, 6, f"=SUM(F{s0}:F{min(s0 + 9, rs - 1)})", FMT_PCT, bold=True, fill=FILL_TOTAL)
    _widths(ws_sup, {"A": 7, "B": 58, "C": 22, "D": 14, "E": 18, "F": 13, "G": 15, "H": 17})
    ws_sup.freeze_panes = "A6"

    # ======================= Sourcing ======================================
    _title(ws_src, f"Sourcing mix - {period.strftime('%B %Y')} and trailing 12 months",
           "Competitive vs non-competitive award value. Shares and totals are formulas.")
    _hdr(ws_src, 4, ["Sourcing type", "This month (CAD)", "T12M (CAD)", "Share of T12M"])
    order = ["Competitive", "Non-competitive", "Advance notice (ACAN)", "Not reported"]
    for i, t in enumerate(order):
        r = 5 + i
        _put(ws_src, r, 1, t)
        _put(ws_src, r, 2, float(src["month_v"].get(t, 0.0)), FMT_USD)
        _put(ws_src, r, 3, float(src["t12_v"].get(t, 0.0)), FMT_USD)
        _put(ws_src, r, 4, f"=IFERROR(C{r}/C$9,0)", FMT_PCT)
    _put(ws_src, 9, 1, "Total", bold=True, fill=FILL_TOTAL)
    _put(ws_src, 9, 2, "=SUM(B5:B8)", FMT_USD, bold=True, fill=FILL_TOTAL)
    _put(ws_src, 9, 3, "=SUM(C5:C8)", FMT_USD, bold=True, fill=FILL_TOTAL)
    _put(ws_src, 9, 4, "=SUM(D5:D8)", FMT_PCT, bold=True, fill=FILL_TOTAL)
    _put(ws_src, 10, 1, "Competitive share of reported value", bold=True)
    _put(ws_src, 10, 2, '=IFERROR(B5/(B5+B6),"n/a")', FMT_PCT, bold=True)
    _put(ws_src, 10, 3, '=IFERROR(C5/(C5+C6),"n/a")', FMT_PCT, bold=True)
    ws_src.cell(row=12, column=1, value="Why contracts were not competed (non-competitive, T12M)").font = F_BOLD
    _hdr(ws_src, 13, ["Reason", "Contracts", "Value (CAD)", "Share of non-competitive"])
    l0 = 14
    for i, row in lim.iterrows():
        r = l0 + i
        _put(ws_src, r, 1, row["reason"][:80])
        _put(ws_src, r, 2, int(row["n"]), FMT_INT)
        _put(ws_src, r, 3, float(row["v"]), FMT_USD)
        _put(ws_src, r, 4, f"=IFERROR(C{r}/C$6,0)", FMT_PCT)
    ws_src.cell(row=l0 + len(lim) + 1, column=1,
                value="Reasons shown are the top 6 by value, so shares may sum to less than 100%.").font = F_NOTE
    _widths(ws_src, {"A": 46, "B": 18, "C": 18, "D": 24})

    # ======================= Amendments ====================================
    _title(ws_am, f"Contract value changes - contracts amended in {period.strftime('%B %Y')}",
           "Stand-alone contracts whose original record is in the file. Change columns are formulas.")
    _put(ws_am, 3, 1, "Contracts with amendment activity this month (any type)", bold=True)
    _put(ws_am, 3, 2, int(amend_n), FMT_INT, bold=True)
    _hdr(ws_am, 5, ["Contract #", "Title", "Supplier", "Category", "Original award (CAD)", "Current value (CAD)",
                    "Change (CAD)", "Change %", "Amendments"])
    a0 = 6
    for i, row in amend.iterrows():
        r = a0 + i
        _put(ws_am, r, 1, row["cn"])
        _put(ws_am, r, 2, row["title"][:70])
        _put(ws_am, r, 3, row["supplier"][:50])
        _put(ws_am, r, 4, row["category"])
        _put(ws_am, r, 5, float(row["orig"]), FMT_USD)
        _put(ws_am, r, 6, float(row["final"]), FMT_USD)
        _put(ws_am, r, 7, f"=F{r}-E{r}", FMT_USD)
        _put(ws_am, r, 8, f'=IFERROR(G{r}/E{r},"n/a")', FMT_PCT)
        _put(ws_am, r, 9, int(row["na"]), FMT_INT)
    if amend.empty:
        ws_am.cell(row=a0, column=1, value="No eligible amended contracts in this month.").font = F_NOTE
    ws_am.cell(row=a0 + max(len(amend), 1) + 1, column=1,
               value="Top 10 by dollar increase. Standing offers and supply arrangements are excluded (their value is a ceiling).").font = F_NOTE
    _widths(ws_am, {"A": 24, "B": 48, "C": 36, "D": 34, "E": 18, "F": 18, "G": 16, "H": 11, "I": 12})
    ws_am.freeze_panes = "A6"

    # ======================= Data Quality ==================================
    _title(ws_dq, "Data quality exceptions (whole file)",
           "Contract-level checks run by the pipeline. Flagged, never silently dropped.")
    _put(ws_dq, 3, 1, "Total contracts in file", bold=True)
    _put(ws_dq, 3, 2, int(n_contracts_total), FMT_INT, bold=True)
    _hdr(ws_dq, 5, ["Issue", "Contracts", "% of all contracts"])
    d0 = 6
    for i, row in dq.iterrows():
        r = d0 + i
        _put(ws_dq, r, 1, row["issue_type"])
        _put(ws_dq, r, 2, int(row["n"]), FMT_INT)
        _put(ws_dq, r, 3, f"=IFERROR(B{r}/$B$3,0)", FMT_PCT)
    _widths(ws_dq, {"A": 56, "B": 14, "C": 18})

    # ======================= Notes ==========================================
    _title(ws_nt, "Notes, definitions and limitations")
    p25, p50, p75 = (float(lag.quantile(x)) for x in (0.25, 0.5, 0.75))
    lines = [
        ("Source", "CanadaBuys contract history, 2023-06-01 onwards (Public Services and Procurement Canada). "
                   "Contains information licensed under the Open Government Licence - Canada."),
        ("Scope", "Federal government contracts. This is public procurement data, not manufacturing or electronics spend."),
        ("Unit of analysis", "One row per contract: amendment history is collapsed to the latest record."),
        ("Spend", "totalContractValue (final cumulative value, CAD). It is identical on every record of a contract, "
                  "so amendments do not double count. Contracts valued at $0 (mostly standing offers and supply "
                  "arrangements) count as contracts but add no spend."),
        ("Award month", "Spend is attributed to the month the contract was first awarded, not when it is paid."),
        ("Late postings", f"Awards are published after the award date: across contracts in this file the median lag is "
                          f"{p50:.0f} days (25th percentile {p25:.0f}, 75th percentile {p75:.0f}). Recent months therefore "
                          f"keep growing as awards are posted. 'Provisional' means the month ended within "
                          f"{C.PROVISIONAL_DAYS} days of the newest publication date - a heuristic, not a guarantee."),
        ("Not available", "The file has no delivery dates, quantities, unit prices or invoices, so on-time delivery, "
                          "price variance and invoice reconciliation cannot be computed from it."),
        ("Categories", "UNSPSC segment mapped to 14 business categories; legacy records without UNSPSC use keyword rules "
                       "on the GSIN description (best effort, not validated)."),
        ("Suppliers", "Rule-based name matching using the publisher's standardized names. No fuzzy matching, so "
                      "related entities (for example subsidiaries) can remain separate."),
        ("Mega-contracts", "A few contracts exceed $1B. Use the 'excl. contracts over $100M' rows for run-rate views."),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M") + " by `python -m spendpipe report`."),
    ]
    for i, (k, v) in enumerate(lines):
        r = 3 + i
        a = ws_nt.cell(row=r, column=1, value=k)
        a.font = F_BOLD
        a.alignment = Alignment(vertical="top")
        b = ws_nt.cell(row=r, column=2, value=v)
        b.font = F_BODY
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws_nt.row_dimensions[r].height = max(18, 15 * (len(v) // 105 + 1))
    _widths(ws_nt, {"A": 18, "B": 110})

    for ws in wb.worksheets:  # print setup: landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = Path(out_path or f"output/monthly_report_{month}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)
